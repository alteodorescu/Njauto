"""
Export the monthly breakdown and yearly rollup to a single Excel file
with two sheets ('Monthly' and 'Yearly'). Reads the CSV produced by
tools/monthly_report.py.
"""

from __future__ import annotations
import csv
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent
CSV_IN = REPO / "reports" / "monthly_report.csv"
XLSX_OUT = REPO / "reports" / "monthly_payouts.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center")
RIGHT = Alignment(horizontal="right")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_monthly():
    rows = []
    with CSV_IN.open() as f:
        for r in csv.DictReader(f):
            rows.append({
                "month": r["month"],
                "lost_overall": int(r["lost_overall"]),
                "lost_daily": int(r["lost_daily"]),
                "lost_total": int(r["lost_total"]),
                "payouts_count": int(r["payouts_count"]),
                "payouts_sum": float(r["payouts_sum_usd"]),
                "challenges_count": int(r["challenges_count"]),
                "fees_sum": float(r["fees_sum_usd"]),
                "profit": float(r["profit_usd"]),
            })
    return rows


def yearly_rollup(monthly):
    yr = defaultdict(lambda: {
        "lost_overall": 0, "lost_daily": 0, "lost_total": 0,
        "payouts_count": 0, "payouts_sum": 0.0,
        "challenges_count": 0, "fees_sum": 0.0, "profit": 0.0,
    })
    for r in monthly:
        y = r["month"][:4]
        yr[y]["lost_overall"] += r["lost_overall"]
        yr[y]["lost_daily"] += r["lost_daily"]
        yr[y]["lost_total"] += r["lost_total"]
        yr[y]["payouts_count"] += r["payouts_count"]
        yr[y]["payouts_sum"] += r["payouts_sum"]
        yr[y]["challenges_count"] += r["challenges_count"]
        yr[y]["fees_sum"] += r["fees_sum"]
        yr[y]["profit"] += r["profit"]
    return [{"year": y, **yr[y]} for y in sorted(yr)]


def style_header(ws, row, col_count):
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def style_total(ws, row, col_count):
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.border = BORDER


def auto_width(ws):
    for col in ws.columns:
        col_letter = col[0].column_letter
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 3
        ws.column_dimensions[col_letter].width = min(max(width, 8), 22)


CURRENCY_COLS = (6, 8, 9)  # Payouts $, Fees $, Profit $

HEADERS_MONTHLY = ["Month", "Lost (overall)", "Lost (daily)", "Lost total",
                   "Payouts #", "Payouts $",
                   "Challenges #", "Fees $", "Profit $"]
HEADERS_YEARLY = ["Year"] + HEADERS_MONTHLY[1:]

# Negative profit gets red font, positive gets green.
RED = Font(color="C00000")
GREEN = Font(color="007A00")


def _apply_styles(ws, last_row, currency_cols, profit_col):
    for row in range(2, last_row + 1):
        for c in currency_cols:
            ws.cell(row=row, column=c).number_format = '"$"#,##0.00'
        for c in range(2, len(HEADERS_MONTHLY) + 1):
            ws.cell(row=row, column=c).alignment = RIGHT
        # Color profit cell by sign.
        v = ws.cell(row=row, column=profit_col).value
        if isinstance(v, (int, float)):
            ws.cell(row=row, column=profit_col).font = GREEN if v >= 0 else RED


def write_monthly(wb, monthly):
    ws = wb.create_sheet("Monthly")
    ws.append(HEADERS_MONTHLY)
    style_header(ws, 1, len(HEADERS_MONTHLY))

    for r in monthly:
        ws.append([
            r["month"], r["lost_overall"], r["lost_daily"], r["lost_total"],
            r["payouts_count"], r["payouts_sum"],
            r["challenges_count"], r["fees_sum"], r["profit"],
        ])

    last = len(monthly) + 1
    ws.append([
        "TOTAL",
        sum(r["lost_overall"] for r in monthly),
        sum(r["lost_daily"] for r in monthly),
        sum(r["lost_total"] for r in monthly),
        sum(r["payouts_count"] for r in monthly),
        sum(r["payouts_sum"] for r in monthly),
        sum(r["challenges_count"] for r in monthly),
        sum(r["fees_sum"] for r in monthly),
        sum(r["profit"] for r in monthly),
    ])
    style_total(ws, last + 1, len(HEADERS_MONTHLY))
    _apply_styles(ws, last + 1, CURRENCY_COLS, profit_col=9)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{last}"
    auto_width(ws)


def write_yearly(wb, yearly):
    ws = wb.create_sheet("Yearly")
    ws.append(HEADERS_YEARLY)
    style_header(ws, 1, len(HEADERS_YEARLY))

    for r in yearly:
        ws.append([
            r["year"], r["lost_overall"], r["lost_daily"], r["lost_total"],
            r["payouts_count"], r["payouts_sum"],
            r["challenges_count"], r["fees_sum"], r["profit"],
        ])

    last = len(yearly) + 1
    ws.append([
        "TOTAL",
        sum(r["lost_overall"] for r in yearly),
        sum(r["lost_daily"] for r in yearly),
        sum(r["lost_total"] for r in yearly),
        sum(r["payouts_count"] for r in yearly),
        sum(r["payouts_sum"] for r in yearly),
        sum(r["challenges_count"] for r in yearly),
        sum(r["fees_sum"] for r in yearly),
        sum(r["profit"] for r in yearly),
    ])
    style_total(ws, last + 1, len(HEADERS_YEARLY))
    _apply_styles(ws, last + 1, CURRENCY_COLS, profit_col=9)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{last}"
    auto_width(ws)


def main():
    monthly = load_monthly()
    yearly = yearly_rollup(monthly)

    wb = openpyxl.Workbook()
    # Remove the default sheet so order is Monthly, Yearly.
    default = wb.active
    wb.remove(default)
    write_monthly(wb, monthly)
    write_yearly(wb, yearly)

    XLSX_OUT.parent.mkdir(exist_ok=True)
    wb.save(XLSX_OUT)
    print(f"Wrote {XLSX_OUT}")
    print(f"  Monthly rows: {len(monthly)}")
    print(f"  Yearly rows:  {len(yearly)}")


if __name__ == "__main__":
    main()
